"""Generador del Reporte Técnico Oficial en formato Word (.docx) para ASHAKids.

Fase: CRUD de Pacientes, Avatares, Gestión Administrativa y Sincronización del Perfil Infantil.
Rama: feat/nlavadom
Repositorio: sromansilva/ashakids-platform
"""

import os
from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

DOCS_DIR = Path(__file__).resolve().parent
OUTPUT_DOCX = DOCS_DIR / "ASHAKids_Reporte_CRUD_Pacientes_y_Perfil_Infantil.docx"

COLOR_PRIMARY = RGBColor(124, 58, 237)    # Violeta (#7C3AED)
COLOR_TEXT = RGBColor(30, 41, 59)         # Slate 800 (#1E293B)
COLOR_MUTED = RGBColor(100, 116, 139)     # Slate 500 (#64748B)
COLOR_HEADER_BG = "7C3AED"
COLOR_ZEBRA_BG = "F8FAFC"
COLOR_BORDER = "E2E8F0"


def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)


def set_cell_margins(cell, top=120, bottom=120, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(
        f'<w:tcMar {nsdecls("w")}>'
        f'<w:top w:w="{top}" w:type="dxa"/>'
        f'<w:bottom w:w="{bottom}" w:type="dxa"/>'
        f'<w:left w:w="{left}" w:type="dxa"/>'
        f'<w:right w:w="{right}" w:type="dxa"/>'
        f'</w:tcMar>'
    )
    tcPr.append(tcMar)


def build_docx_report():
    doc = Document()

    # Configuración de márgenes
    for section in doc.sections:
        section.top_margin = Inches(0.9)
        section.bottom_margin = Inches(0.9)
        section.left_margin = Inches(0.9)
        section.right_margin = Inches(0.9)

    # ----------------------------------------------------
    # PORTADA Y ENCABEZADO
    # ----------------------------------------------------
    p_title = doc.add_paragraph()
    p_title.paragraph_format.space_before = Pt(0)
    p_title.paragraph_format.space_after = Pt(2)
    run_t = p_title.add_run("ASHAKids Platform")
    run_t.font.name = "Arial"
    run_t.font.size = Pt(24)
    run_t.font.bold = True
    run_t.font.color.rgb = COLOR_PRIMARY

    p_sub = doc.add_paragraph()
    p_sub.paragraph_format.space_after = Pt(8)
    run_sub = p_sub.add_run("Reporte Técnico — CRUD de Pacientes, Avatares, Gestión Administrativa y Perfil Infantil")
    run_sub.font.name = "Arial"
    run_sub.font.size = Pt(13)
    run_sub.font.bold = True
    run_sub.font.color.rgb = RGBColor(15, 23, 42)

    p_meta = doc.add_paragraph()
    p_meta.paragraph_format.space_after = Pt(16)
    run_meta = p_meta.add_run(
        "Rama: feat/nlavadom  |  Fecha: Octubre 2026  |  Arquitectura: React + Vite + FastAPI + Supabase PostgreSQL\n"
        "Ámbitos: padre/config (Mis hijos), admin/cuentas (Ver/Reactivar hijos), Centro Familiar, Mi Camino ASHA"
    )
    run_meta.font.name = "Arial"
    run_meta.font.size = Pt(9.5)
    run_meta.font.italic = True
    run_meta.font.color.rgb = COLOR_MUTED

    # ----------------------------------------------------
    # SECCIONES PRINCIPALES
    # ----------------------------------------------------
    content = [
        ("1. Resumen Ejecutivo", [
            "El presente informe documenta la implementación completa del ciclo de vida y gestión de pacientes (hijos asociados a sus padres o tutores) en la plataforma ASHAKids, respetando estrictamente la arquitectura React → FastAPI → Supabase PostgreSQL.",
            "Principales hitos alcanzados:",
            "• CRUD Completo de Pacientes para el Padre: Creación atómica (paciente + perfil 1:1), consulta, edición autorizada con validación IDOR, e inactivación lógica.",
            "• Gestión Administrativa en admin/cuentas: Visualización de todos los hijos asociados a un tutor (activos e inactivos) y capacidad de reactivación restringida exclusivamente al rol ADMIN.",
            "• Persistencia de Avatares: Utilización de la columna nativa PACIENTES.avatar_nombre con validación de catálogo (10 avatares de animales), valor predeterminado 'zorro' y fallback visual.",
            "• Preparación de Assets en Frontend: Estructura frontend/public/avatars/ con 10 archivos SVG vectoriales, catálogo centralizado y documentación README.md de extensibilidad.",
            "• Protección Estricta del Historial Clínico: Bloqueo de eliminación física con código HTTP 409 Conflict si existen dependencias médicas o históricas asociadas.",
            "• Auditoría Centralizada: Registro completo de cada operación (INSERT, UPDATE, DELETE) en la tabla public.auditoria_cambios sin comprometer datos sensibles.",
            "• Sincronización Global del Perfil Infantil: Implementación de ChildContext y selector unificado en la cabecera y barra lateral del panel de padres, conectando datos reales en Centro Familiar y Mi Camino ASHA."
        ]),

        ("2. Restricciones y Arquitectura de Base de Datos", [
            "Conforme a las directrices de la auditoría y diseño previo:",
            "1. Columna de Avatar Existente:",
            "   - Se verificó que la columna PACIENTES.avatar_nombre VARCHAR(20) NOT NULL DEFAULT 'zorro' ya existía en Supabase PostgreSQL.",
            "   - No se ejecutaron migraciones DDL redundantes, no se crearon tablas auxiliares ni se alteró el tipo de dato.",
            "   - Almacena el identificador nemotécnico (slug) del avatar ('zorro', 'oso', 'conejo', 'panda', 'leon', 'koala', 'tortuga', 'buho', 'mono', 'pinguino'), nunca rutas directas.",
            "2. Columna Sexo:",
            "   - PACIENTES.sexo VARCHAR(10) NOT NULL validado estrictamente en Backend y Frontend contra las opciones: ['Masculino', 'Femenino', 'Otro'].",
            "3. Integridad Referencial y Perfiles 1:1:",
            "   - PACIENTES.id_tutor ForeignKey a tutores.id_tutor (relación 1:N).",
            "   - PERFILES.id_paciente ForeignKey único a pacientes.id_paciente (relación 1:1).",
            "   - La creación de Paciente y Perfil se ejecuta de forma atómica en una única transacción de SQLAlchemy. Si la inserción del perfil falla, la transacción se revierte íntegramente.",
            "4. Preservación del Historial Clínico y Dependencias:",
            "   - Se identificaron claves foráneas en tablas clínicas (expedientes, tratamientos, sesiones, reservas, evaluaciones, resultados_actividad, asignaciones_actividades).",
            "   - El Backend verifica explícitamente estas tablas antes de permitir un borrado físico, rechazando con 409 Conflict si existen datos vinculados."
        ]),

        ("3. Preparación de la Carpeta de Avatares en Frontend", [
            "Ubicación oficial: frontend/public/avatars/",
            "• Se incorporaron 10 avatares SVG estilizados: zorro.svg, oso.svg, conejo.svg, panda.svg, leon.svg, koala.svg, tortuga.svg, buho.svg, mono.svg, pinguino.svg.",
            "• Se añadió frontend/public/avatars/README.md detallando:",
            "   a) Convención de nombres: formato SVG vectorial cuadrado (relación 1:1, viewbox 0 0 100 100), optimizado para rendimiento web.",
            "   b) Correspondencia con PACIENTES.avatar_nombre: el nombre base del archivo sin extensión coincide exactamente con el identificador del catálogo.",
            "   c) Procedimiento para incorporar nuevos avatares: 1) Añadir el archivo SVG en public/avatars/, 2) Registrar el slug en AVATAR_CATALOG de frontend/src/config/avatars.ts, 3) Añadir el slug en ALLOWED_AVATARS de backend/app/schemas/pacientes.py. Sin necesidad de tocar la base de datos.",
            "• Catálogo Unificado (frontend/src/config/avatars.ts):",
            "   - Exporta AVATAR_CATALOG, DEFAULT_AVATAR_ID ('zorro') y helper getAvatarInfo(slug).",
            "   - Provee fallback seguro hacia 'zorro' si se encuentra un identificador desconocido o la imagen no carga."
        ]),

        ("4. CRUD de Pacientes para el Padre (padre/config → Mis Hijos)", [
            "Flujos implementados y validados:",
            "• Listar hijos: Endpoint GET /api/v1/padres/hijos resuelve el id_tutor desde la sesión del usuario autenticado (id_usuario). Filtra automáticamente activo = TRUE, excluyendo hijos inactivos.",
            "• Crear hijo: Modal reactivo con campos requeridos: Nombres, Apellidos, Fecha de nacimiento (con control de fecha futura y rango de edad pediátrica), Sexo (selector Masculino/Femenino/Otro) y selector visual en cuadrícula de los 10 avatares disponibles. Creación atómica de Paciente + Perfil infantil.",
            "• Editar hijo: Modal reactivo que permite actualizar nombres, apellidos, fecha de nacimiento, sexo y avatar. Verificación rigurosa de pertenencia al tutor autenticado (evita IDOR con 404).",
            "• Inactivar hijo: Operación lógica mediante PATCH /api/v1/padres/hijos/{id}/inactivar. Cambia activo = FALSE, preservando intactos el perfil, sesiones y registros clínicos. Si era el hijo seleccionado actualmente, el ChildContext transiciona automáticamente al siguiente hijo activo o a null.",
            "• Eliminación física condicionada: DELETE /api/v1/padres/hijos/{id}. El servicio audita todas las dependencias clínicas. Si existen registros dependientes, aborta inmediatamente con HTTP 409 Conflict y mensaje descriptivo."
        ]),

        ("5. Gestión Administrativa de Hijos (admin/cuentas)", [
            "Flujos implementados en el panel de Administración:",
            "• En la vista de detalle de cuentas de rol 'PADRE', se agregó el botón 'Ver hijos' con icono descriptivo.",
            "• Modal Administrativo 'Hijos del Tutor':",
            "   - Muestra el listado completo de hijos asociados al tutor consultado vía GET /api/v1/admin/cuentas/{id_usuario}/hijos.",
            "   - Despliega avatar, nombre completo, edad, sexo, fecha de nacimiento y estado (Activo / Inactivo).",
            "• Reactivación Administrativa:",
            "   - Para pacientes con estado inactivo, se expone la acción 'Reactivar'.",
            "   - Endpoint protegido PATCH /api/v1/admin/pacientes/{id_paciente}/reactivar.",
            "   - Requiere autenticación con rol ADMIN (require_admin). Usuarios con rol PADRE o TERAPEUTA reciben HTTP 403 Forbidden.",
            "   - Restablece activo = TRUE, conserva todos los datos históricos y deja rastro de auditoría con el id_usuario del administrador.",
            "   - Tras reactivarse, el hijo reaparece de inmediato en los selectores y vistas del padre."
        ]),

        ("6. Auditoría Centralizada (public.auditoria_cambios)", [
            "Todas las mutaciones sobre la entidad Paciente se registran mediante el servicio central auditoria_service.py:",
            "• nombre_tabla: 'pacientes'",
            "• nombre_entidad: 'Paciente'",
            "• accion: 'INSERT' (creación), 'UPDATE' (edición e inactivación/reactivación), 'DELETE' (borrado físico)",
            "• datos_anteriores: Snapshot JSON del registro previo a la modificación",
            "• datos_nuevos: Snapshot JSON del nuevo estado tras la mutación",
            "• id_usuario_actor: ID del usuario autenticado que ejecutó la acción (padre o administrador)",
            "• fecha_evento: Timestamp UTC actual",
            "Seguridad de la información: No se incluyen contraseñas, tokens de sesión ni credenciales en las trazas de auditoría."
        ]),

        ("7. Sincronización Global del Perfil Infantil (Centro Familiar y Mi Camino ASHA)", [
            "Para unificar la experiencia del padre y evitar inconsistencias entre vistas:",
            "• Contexto Global (frontend/src/context/ChildContext.tsx):",
            "   - Administra el estado global del hijo seleccionado: activeChild, childrenList, isLoading, selectChild, refreshChildren.",
            "   - Persiste la selección en localStorage ('ashakids_selected_child_id') para preservar el contexto ante recargas de página (F5).",
            "   - Descarta automáticamente hijos que hayan sido inactivados.",
            "• Componentes Actualizados:",
            "   - Barra Lateral (PadreSidebar): Renderiza los hijos reales con su avatar respectivo, edad calculada e indicador visual de hijo activo.",
            "   - Selector Superior (ChildPicker): Selector interactivo para alternar entre hijos activos desde cualquier pantalla del área de padres.",
            "   - Centro Familiar (PadreHome): Saludo personalizado y datos del hijo seleccionado activo.",
            "   - Mi Camino ASHA (MiCaminoAsha): Sincronizado con activeChild para el encabezado y la exportación de reportes PDF.",
            "• Clasificación de Datos (Reales vs Simulados):",
            "   - Datos Reales Persistidos: Identidad del niño, nombres, apellidos, edad, sexo, fecha de nacimiento y avatar.",
            "   - Datos Pendientes / Simulados: Objetivos clínicos detallados, porcentajes de avance acumulado y rachas globales.",
            "   - Justificación Técnica: La tabla OBJETIVOS posee la columna resultados_objetivo como JSONB sin estructura clínica consolidada en esta fase. No se inventaron métricas artificiales; las evaluaciones clínicas y el rendimiento lúdico se mantienen desacoplados hasta que se defina la estructura formal de resultados."
        ]),

        ("8. Validación, Compilación y Pruebas Automatizadas", [
            "Se implementaron y ejecutaron suites de prueba tanto en backend como en frontend:",
            "1. Backend (backend/tests/test_pacientes_crud.py):",
            "   - 13 pruebas unitarias y de integración end-to-end con conexión real a PostgreSQL.",
            "   - test_01: Creación atómica de paciente y perfil 1:1 con avatar 'panda'. (APROBADO)",
            "   - test_02: Asignación por defecto de avatar 'zorro' ante omisión. (APROBADO)",
            "   - test_03: Rechazo HTTP 422 ante identificador de avatar fuera del catálogo. (APROBADO)",
            "   - test_04: Validación estricta de sexo ('Masculino', 'Femenino', 'Otro') y rechazo de valores arbitrarios. (APROBADO)",
            "   - test_05: Rechazo HTTP 422 de fechas de nacimiento futuras. (APROBADO)",
            "   - test_06: Protección IDOR entre tutores: un padre no puede ver, editar, inactivar ni eliminar hijos de otro tutor. (APROBADO)",
            "   - test_07: Edición autorizada de datos y cambio de avatar a 'buho'. (APROBADO)",
            "   - test_08: Inactivación lógica: activo=False y exclusión de los listados normales del padre. (APROBADO)",
            "   - test_09: Bloqueo de eliminación física con HTTP 409 Conflict ante registros clínicos dependientes. (APROBADO)",
            "   - test_10: Eliminación física admisible cuando el paciente no tiene dependencias clínicas. (APROBADO)",
            "   - test_11: Vista administrativa: consulta de hijos activos e inactivos de un tutor. (APROBADO)",
            "   - test_12: Reactivación administrativa: éxito con rol ADMIN y bloqueo HTTP 403 para usuarios con rol PADRE. (APROBADO)",
            "   - test_13: Verificación de persistencia en tabla central AUDITORIA_CAMBIOS. (APROBADO)",
            "2. Frontend (Compilación y TypeScript):",
            "   - npm run build ejecutado con éxito total (0 errores de TypeScript, tiempo: 7.64s).",
            "   - Integración completa con Vite y React Router v7 sin regresiones."
        ]),

        ("9. Inventario de Archivos Modificados y Añadidos", [
            "Archivos de Backend Modificados/Añadidos:",
            "• backend/app/models/perfiles.py: Declaración formal de Paciente con columna avatar_nombre y relación id_tutor ForeignKey.",
            "• backend/app/schemas/pacientes.py: Esquemas Pydantic v2 (CrearHijoRequest, ActualizarHijoRequest, PerfilInfantilResponse, HijoItemResponse, HijoListResponse, OperacionHijoResponse) con validación de catálogo ALLOWED_AVATARS y ALLOWED_SEXO.",
            "• backend/app/services/pacientes_service.py: Lógica de negocio, transacciones atómicas, guardias IDOR, exclusión de inactivos, verificación de dependencias clínicas para 409 Conflict y registro en auditoria_cambios.",
            "• backend/app/api/v1/padres.py: Endpoints REST para el padre (/hijos, /hijos/{id}, /hijos/{id}/inactivar, /hijos/{id}).",
            "• backend/app/api/v1/admin.py: Endpoints administrativos (/cuentas/{id}/hijos, /pacientes/{id}/reactivar).",
            "• backend/tests/test_pacientes_crud.py: Suite de 13 pruebas automatizadas con cobertura completa de casos de uso y seguridad.",
            "",
            "Archivos de Frontend Modificados/Añadidos:",
            "• frontend/public/avatars/ (10 archivos SVG): zorro.svg, oso.svg, conejo.svg, panda.svg, leon.svg, koala.svg, tortuga.svg, buho.svg, mono.svg, pinguino.svg.",
            "• frontend/public/avatars/README.md: Guía técnica de assets y catálogo de avatares.",
            "• frontend/src/config/avatars.ts: Catálogo centralizado, valores por defecto y funciones de fallback.",
            "• frontend/src/types/pacientes.ts: Definiciones de interfaces TypeScript del dominio infantil.",
            "• frontend/src/services/pacientesService.ts: Cliente de API para operaciones de pacientes del padre.",
            "• frontend/src/services/adminService.ts: Métodos administrativos getHijosDePadre y reactivarHijo.",
            "• frontend/src/context/ChildContext.tsx: Proveedor global de sincronización infantil y selección activa.",
            "• frontend/src/main.tsx: Envoltura de la aplicación en ChildProvider.",
            "• frontend/src/App.tsx: Integración de PadreSidebar, ChildPicker, PadreHome y PadreConfig (pestaña 'hijos' con modales de creación, edición, inactivación y eliminación).",
            "• frontend/src/pages/admin/AdminCuentas.tsx: Botón 'Ver hijos', modal con listado de activos/inactivos y acción 'Reactivar'.",
            "• frontend/src/pages/terapeuta/MiCaminoAsha.tsx: Sincronización del paciente activo en visualización y reportes PDF."
        ])
    ]

    # Renderizado de secciones y párrafos
    for sec_title, sec_paragraphs in content:
        h = doc.add_heading(sec_title, level=2)
        h.paragraph_format.space_before = Pt(14)
        h.paragraph_format.space_after = Pt(6)

        for p_text in sec_paragraphs:
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(4)
            p.paragraph_format.line_spacing = 1.15
            run = p.add_run(p_text)
            run.font.name = "Arial"
            run.font.size = Pt(10)
            if p_text.startswith("•") or p_text.startswith("   -"):
                run.font.color.rgb = RGBColor(51, 65, 85)
            else:
                run.font.color.rgb = COLOR_TEXT

    # Guardar documento
    doc.save(OUTPUT_DOCX)
    print(f"Documento Word generado exitosamente en: {OUTPUT_DOCX}")


if __name__ == "__main__":
    build_docx_report()
