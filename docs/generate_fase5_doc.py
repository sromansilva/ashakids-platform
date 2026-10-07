"""Generador del Reporte Word para ASHAKids - Fase de Navegación, Routing URL y Protección por Roles (Fase 5)."""

import os
from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

DOCS_DIR = Path(__file__).resolve().parent
OUTPUT_DOCX = DOCS_DIR / "ASHAKids_Reporte_Fase_5_Navegacion_y_Routing.docx"


def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)


def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(
        f'<w:tcMar {nsdecls("w")}>'
        f'<w:top w:w="{top}" w:type="dxa"/>'
        f'<w:bottom w:w="{bottom}" w:type="dxa"/>'
        f'<w:left w:w="{left}" w:type="dxa"/>'
        f'<w:right w:w="{right}" w:type="dxa"/>'
        f'</w:tcMar>'
    )
    tcPr.append(tcMar)


def build_docx_report():
    doc = Document()

    for section in doc.sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)

    # Portada / Título
    title_p = doc.add_paragraph()
    title_p.paragraph_format.space_before = Pt(0)
    title_p.paragraph_format.space_after = Pt(4)
    run_t = title_p.add_run("ASHAKids Platform")
    run_t.font.name = "Arial"
    run_t.font.size = Pt(24)
    run_t.font.bold = True
    run_t.font.color.rgb = RGBColor(124, 58, 237)  # Violeta

    sub_p = doc.add_paragraph()
    sub_p.paragraph_format.space_after = Pt(16)
    run_sub = sub_p.add_run("Reporte Técnico — Fase 5: Navegación, Routing URL Real y Protección por Roles")
    run_sub.font.name = "Arial"
    run_sub.font.size = Pt(14)
    run_sub.font.bold = True
    run_sub.font.color.rgb = RGBColor(28, 17, 53)

    meta_p = doc.add_paragraph()
    meta_p.paragraph_format.space_after = Pt(20)
    run_meta = meta_p.add_run("Rama: feat/nlavadom | Fecha: Octubre 2026 | Arquitectura: React + Vite + React Router v7 + FastAPI")
    run_meta.font.name = "Arial"
    run_meta.font.size = Pt(10)
    run_meta.font.italic = True
    run_meta.font.color.rgb = RGBColor(107, 94, 138)

    # Secciones
    sections = [
        ("1. Estado Inicial del Routing", [
            "Antes de la presente fase, la aplicación utilizaba un estado local reactivo en App.tsx (const [view, setView]) acoplado a un despachador interno go(targetView) y un switch de renderizado renderView().",
            "La URL del navegador se mantenía estática en http://localhost:5173/ sin importar la navegación realizada.",
            "Al presionar F5 o recargar el navegador, se perdía la vista actual.",
            "No existía soporte para deep linking ni navegación mediante el historial del navegador (atrás/adelante).",
            "La protección entre roles no estaba vinculada canónicamente a la ruta solicitada."
        ]),
        ("2. Librería de Routing Utilizada", [
            "Se implementó react-router-dom v7.13.0, alineada con el paquete base react-router ya existente en el proyecto.",
            "En frontend/src/main.tsx, se integró <BrowserRouter> envolviendo al <AuthProvider> existente y al componente raíz <App />.",
            "Se mantuvo la reactividad y el ciclo de vida sin provocar renderizados dobles ni recargas de página."
        ]),
        ("3. Mapeo y Rutas Canónicas Implementadas", [
            "Se creó el módulo frontend/src/routes/paths.ts con tres utilidades principales:",
            "• viewToPath(view): Convierte identificadores internos (ej. 'public/especialistas', 'padre/hijos') a URLs canónicas.",
            "• pathToView(pathname, userRole): Traduce rutas de URL a identificadores internos de vista.",
            "• getRequiredRoleForPath(pathname, userRole): Identifica qué rol semántico (PADRE, TERAPEUTA, ADMIN) es requerido.",
            "Rutas públicas: /, /login, /especialistas, /especialidades, /recursos, /sobre-nosotros, /planes, /ayuda, /contacto, /trabaja, /ashi, /historias.",
            "Área Padre: /padre, /padre/pacientes, /padre/agenda, /padre/psicologos, /mundo-asha/*, /session/*, /pay/*.",
            "Área Terapeuta: /terapeuta, /terapeuta/agenda, /terapeuta/pacientes, /terapeuta/reportes, /terapeuta/analiticas, etc.",
            "Área Admin: /admin, /admin/cuentas, /admin/terapeutas, /admin/operacion, /admin/pagos, /admin/ml, /admin/auditoria, etc."
        ]),
        ("4. Integración con AuthContext y Sesión HttpOnly", [
            "Se preservó el sistema de sesión basado en cookies seguras HttpOnly (ashakids_session).",
            "AuthContext revalida la sesión contra GET /api/v1/auth/me en el montaje inicial.",
            "Durante el período de verificación (isLoading === true), las rutas protegidas despliegan una pantalla de carga segura con spinner animado, evitando rebotes erróneos a /login.",
            "Al cerrar sesión con handleLogout(), se invalida la sesión en backend mediante POST /api/v1/auth/logout y se navega a /login."
        ]),
        ("5. ProtectedRoute y RoleRoute", [
            "Se reforzó la seguridad en dos niveles:",
            "1. Frontend: ProtectedRoute y RoleRoute impiden el renderizado no autorizado.",
            "   - Si un usuario sin sesión intenta entrar a /padre, /terapeuta o /admin, es redirigido a /login.",
            "   - Si un usuario autenticado intenta acceder al área de otro rol (ej. PADRE en /admin), se renderiza la pantalla estilizada de Acceso Restringido con botón para volver a su panel correspondiente.",
            "2. Backend: Los endpoints FastAPI validan sesión y rol mediante require_padre, require_terapeuta y require_admin, retornando 401 o 403 según corresponda. La URL jamás determina la seguridad en el servidor."
        ]),
        ("6. Transición del Sistema view / go / renderView", [
            "La URL del navegador es ahora la única fuente de verdad: const view: View = pathToView(location.pathname, authRole).",
            "El callback go(target) fue adaptado para ejecutar navigate(viewToPath(target)), logrando que los más de 24 submódulos sigan funcionando idénticamente sin alterar sus callbacks internos.",
            "DashLayout recibe la vista activa y mantiene resaltadas las pestañas correspondientes en tiempo real."
        ]),
        ("7. Resolución de Recargas (F5) y SPA Fallback", [
            "Vite provee soporte nativo para SPA History Fallback en modo de desarrollo.",
            "Las solicitudes directas a /padre, /padre/pacientes, /terapeuta y /admin retornan código HTTP 200 y el index.html.",
            "Al hacer F5, React Router detecta la URL actual, AuthContext revalida la cookie de sesión y la vista se restaura sin pantallas en blanco."
        ]),
        ("8. Pruebas y Validación", [
            "• Suite Automatizada de Frontend (frontend/tests/routing.test.mjs): 25 casos de prueba ejecutados y aprobados (100% de éxito).",
            "• Suite de Pruebas de Backend (backend/tests): 40 casos de prueba ejecutados y aprobados con conexión real a Supabase PostgreSQL.",
            "• Compilación de Producción: npm run build ejecutado exitosamente sin errores de TypeScript ni empaquetado (7.66s)."
        ])
    ]

    for sec_title, sec_paragraphs in sections:
        h = doc.add_heading(sec_title, level=2)
        h.paragraph_format.space_before = Pt(14)
        h.paragraph_format.space_after = Pt(6)
        for p_text in sec_paragraphs:
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(4)
            p.paragraph_format.line_spacing = 1.15
            run = p.add_run(p_text)
            run.font.name = "Arial"
            run.font.size = Pt(10)
            run.font.color.rgb = RGBColor(40, 40, 40)

    doc.save(OUTPUT_DOCX)
    print(f"Reporte Word generado exitosamente en: {OUTPUT_DOCX}")


if __name__ == "__main__":
    build_docx_report()
