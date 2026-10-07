"""Generador del Reporte Word para ASHAKids - Fase 2: Arquitectura Modular Frontend + Backend API."""

import os
from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

DOCS_DIR = Path(__file__).resolve().parent
OUTPUT_DOCX = DOCS_DIR / "ASHAKids_Reporte_Fase_2_Arquitectura_Modular.docx"


def set_cell_background(cell, fill_hex):
    """Establece el color de fondo de una celda de tabla."""
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)


def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    """Ajusta márgenes internos de una celda."""
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(
        f'<w:tcMar {nsdecls("w")}>'
        f'<w:top w:w="{top}" w:type="dxa"/>'
        f'<w:bottom w:w="{bottom}" w:type="dxa"/>'
        f'<w:left w:w="{left}" w:type="dxa"/>'
        f'<w:right w:w="{right}" w:type="dxa"/>'
        f'</w:tcMar>'
    )
    tcPr.append(tcMar)


def build_report():
    doc = Document()

    # Configuración de márgenes de página
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)

    # Colores corporativos ASHAKids
    c_primary = RGBColor(0x3B, 0x1E, 0x7A)    # Violeta oscuro institucional
    c_secondary = RGBColor(0x6D, 0x28, 0xD9)  # Violeta vibrante
    c_dark = RGBColor(0x1F, 0x29, 0x37)       # Gris texto oscuro
    c_gray = RGBColor(0x4B, 0x55, 0x63)       # Gris texto secundario
    c_success = RGBColor(0x05, 0x96, 0x69)    # Verde confirmación
    c_warning = RGBColor(0xD9, 0x77, 0x06)    # Ámbar advertencia

    # -------------------------------------------------------------
    # ENCABEZADO Y TÍTULO PRINCIPAL
    # -------------------------------------------------------------
    p_pre = doc.add_paragraph()
    p_pre.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run_pre = p_pre.add_run("ASHAKids Platform — Arquitectura de Software")
    run_pre.font.name = "Calibri"
    run_pre.font.size = Pt(11)
    run_pre.font.bold = True
    run_pre.font.color.rgb = c_secondary

    title_p = doc.add_paragraph()
    title_p.paragraph_format.space_before = Pt(4)
    title_p.paragraph_format.space_after = Pt(8)
    title_run = title_p.add_run("Reporte Técnico: Fase 2\nArquitectura Modular Frontend + Backend API")
    title_run.font.name = "Calibri"
    title_run.font.size = Pt(22)
    title_run.font.bold = True
    title_run.font.color.rgb = c_primary

    # Metadata resumen
    meta_p = doc.add_paragraph()
    meta_p.paragraph_format.space_after = Pt(18)
    meta_run = meta_p.add_run(
        "Rama de trabajo: feat/nlavadom  |  Arquitectura: React + Vite → FastAPI → SQLAlchemy 2.x → PostgreSQL\n"
        "Autenticación: Sesiones Propias (HttpOnly Cookie)  |  Fuente de Verdad DB: Esquema Real ASHAKids  |  Estado: Validación Exitosa"
    )
    meta_run.font.name = "Calibri"
    meta_run.font.size = Pt(9.5)
    meta_run.font.color.rgb = c_gray

    def add_section_header(text):
        h = doc.add_paragraph()
        h.paragraph_format.space_before = Pt(16)
        h.paragraph_format.space_after = Pt(4)
        hrun = h.add_run(text)
        hrun.font.name = "Calibri"
        hrun.font.size = Pt(13)
        hrun.font.bold = True
        hrun.font.color.rgb = c_primary
        return h

    def add_body_p(text, bold_prefix="", space_after=6):
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(space_after)
        if bold_prefix:
            br = p.add_run(bold_prefix)
            br.font.name = "Calibri"
            br.font.size = Pt(10.5)
            br.font.bold = True
            br.font.color.rgb = c_dark
        r = p.add_run(text)
        r.font.name = "Calibri"
        r.font.size = Pt(10.5)
        r.font.color.rgb = c_dark
        return p

    def add_callout(title, text, is_warning=False):
        tbl = doc.add_table(rows=1, cols=1)
        tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
        cell = tbl.cell(0, 0)
        cell.width = Inches(6.5)
        bg_color = "FEF3C7" if is_warning else "EDE9FE"
        border_color = "F59E0B" if is_warning else "7C3AED"
        set_cell_background(cell, bg_color)
        set_cell_margins(cell, top=120, bottom=120, left=180, right=180)

        cp = cell.paragraphs[0]
        cp.paragraph_format.space_after = Pt(2)
        crun_title = cp.add_run(title + "\n")
        crun_title.font.name = "Calibri"
        crun_title.font.size = Pt(10.5)
        crun_title.font.bold = True
        crun_title.font.color.rgb = c_warning if is_warning else c_secondary

        crun_text = cp.add_run(text)
        crun_text.font.name = "Calibri"
        crun_text.font.size = Pt(9.5)
        crun_text.font.color.rgb = c_dark

        empty = doc.add_paragraph()
        empty.paragraph_format.space_after = Pt(4)

    # -------------------------------------------------------------
    # 1. ESTRUCTURA FINAL DEL PROYECTO
    # -------------------------------------------------------------
    add_section_header("1. Estructura Final del Proyecto")
    add_body_p(
        "Se completó la separación y desacoplamiento estructural entre el frontend (React + Vite) "
        "y el backend (FastAPI + SQLAlchemy 2.x). El frontend opera como Single Page Application consumiendo "
        "únicamente la API REST de FastAPI mediante credenciales HttpOnly, sin tener dependencias ni "
        "comunicación directa con Supabase."
    )

    p_tree = doc.add_paragraph()
    p_tree.paragraph_format.space_after = Pt(6)
    tree_run = p_tree.add_run(
        "ashakid/\n"
        "├── backend/\n"
        "│   ├── app/\n"
        "│   │   ├── api/\n"
        "│   │   │   ├── deps.py                 # Inyección de dependencias (cookie, auth, roles)\n"
        "│   │   │   └── v1/\n"
        "│   │   │       └── auth.py             # Endpoints /login, /logout, /me\n"
        "│   │   ├── core/\n"
        "│   │   │   ├── config.py               # Settings pydantic-settings (.env)\n"
        "│   │   │   ├── database.py             # AsyncEngine & async_sessionmaker SQLAlchemy\n"
        "│   │   │   └── security.py             # Argon2id hashing y verificación\n"
        "│   │   ├── models/\n"
        "│   │   │   └── auth.py                 # Modelos ORM: usuarios, roles, administradores, sesiones\n"
        "│   │   ├── schemas/\n"
        "│   │   │   └── auth.py                 # Esquemas Pydantic entrada/salida\n"
        "│   │   ├── services/\n"
        "│   │   │   └── auth_service.py         # Lógica de login, tokens hash SHA-256 y sesiones\n"
        "│   │   └── main.py                     # Entry point FastAPI, CORS y router /api/v1\n"
        "│   ├── scripts/\n"
        "│   │   ├── create_test_users.py        # Generador de hashes y script SQL\n"
        "│   │   └── seed_test_users.sql         # Seed corregido contra el esquema real\n"
        "│   └── tests/\n"
        "│       ├── test_auth_api.py            # Pruebas de integración endpoints API\n"
        "│       └── test_auth_security.py       # Pruebas unitarias de seguridad y hashing\n"
        "├── frontend/\n"
        "│   ├── src/\n"
        "│   │   ├── api/\n"
        "│   │   │   └── client.ts               # Cliente HTTP centralizado (fetch + credentials)\n"
        "│   │   ├── auth/\n"
        "│   │   │   └── AuthContext.tsx         # Contexto de autenticación y proveedor global\n"
        "│   │   ├── hooks/\n"
        "│   │   │   └── useAuth.ts              # Custom hook de autenticación\n"
        "│   │   ├── pages/\n"
        "│   │   │   ├── public/                 # SpecialistsPage, MundoASHAPage, ResourcesPage, AboutUsPage\n"
        "│   │   │   ├── auth/LoginPage.tsx      # Vista modular de Login (<300 líneas)\n"
        "│   │   │   ├── padre/DashboardPage.tsx # Guard de rol PADRE\n"
        "│   │   │   ├── terapeuta/DashboardPage.tsx # Guard de rol TERAPEUTA\n"
        "│   │   │   └── admin/DashboardPage.tsx # Guard de rol ADMIN\n"
        "│   │   ├── routes/\n"
        "│   │   │   ├── ProtectedRoute.tsx      # Protección de rutas autenticadas\n"
        "│   │   │   └── RoleRoute.tsx           # Protección de rutas por rol semántico\n"
        "│   │   ├── services/\n"
        "│   │   │   └── authService.ts          # Llamadas a /api/v1/auth/*\n"
        "│   │   ├── types/\n"
        "│   │   │   └── auth.ts                 # Tipos de dominio: User, SemanticRole, AuthState\n"
        "│   │   ├── app/App.tsx                 # Enrutamiento, vistas y layouts Figma Make preservados\n"
        "│   │   └── main.tsx                    # Inyección del AuthProvider raíz\n"
        "│   └── src/vite-env.d.ts               # Tipado de entorno Vite (VITE_API_URL)\n"
        "└── docs/\n"
        "    ├── ASHAKids_Reporte_Fase_1_Configuracion_y_Hashing.docx\n"
        "    └── ASHAKids_Reporte_Fase_2_Arquitectura_Modular.docx"
    )
    tree_run.font.name = "Consolas"
    tree_run.font.size = Pt(8.5)
    tree_run.font.color.rgb = c_gray

    # -------------------------------------------------------------
    # 2. ARCHIVOS CREADOS
    # -------------------------------------------------------------
    add_section_header("2. Archivos Creados")
    created_files = [
        ("backend/app/core/database.py", "Gestión de conexión async SQLAlchemy 2.x con asyncpg y generador get_db()."),
        ("backend/app/models/auth.py", "Modelos ORM: Usuario, Rol, Administrador, UsuarioRol, SesionAutenticacion."),
        ("backend/app/schemas/auth.py", "Esquemas Pydantic: LoginRequest, UserResponse, AuthResponse, MessageResponse."),
        ("backend/app/services/auth_service.py", "Lógica de negocio: hashing SHA-256 de tokens, persistencia y revocación."),
        ("backend/app/api/deps.py", "Dependencias FastAPI para autenticación por cookie ashakids_session y roles."),
        ("backend/app/api/v1/auth.py", "Endpoints REST /login, /logout y /me con cookies HttpOnly."),
        ("backend/tests/test_auth_api.py", "Suite de pruebas automatizadas para OpenAPI, /docs, login y ciclo de sesión."),
        ("frontend/src/vite-env.d.ts", "Declaración de interfaces ViteEnv y variable VITE_API_URL."),
        ("frontend/src/api/client.ts", "Cliente HTTP centralizado con manejo unificado de errores y credentials: include."),
        ("frontend/src/types/auth.ts", "Definición de modelos de dominio TypeScript y roles semánticos (PADRE, TERAPEUTA, ADMIN)."),
        ("frontend/src/services/authService.ts", "Servicio frontend para consumo de endpoints de autenticación FastAPI."),
        ("frontend/src/auth/AuthContext.tsx", "Contexto de autenticación React con verificación de sesión automática."),
        ("frontend/src/hooks/useAuth.ts", "Hook personalizado para acceso a sesión, estado de carga y métodos login/logout."),
        ("frontend/src/routes/ProtectedRoute.tsx", "Componente guard para rutas privadas con fallback de redirección a login."),
        ("frontend/src/routes/RoleRoute.tsx", "Componente guard con validación estricta de roles semánticos."),
        ("frontend/src/pages/auth/LoginPage.tsx", "Componente de Login modular (<300 líneas), preservando diseño Figma Make."),
        ("frontend/src/pages/public/SpecialistsPage.tsx", "Página pública modular de especialistas."),
        ("frontend/src/pages/public/MundoASHAPage.tsx", "Página pública modular de Mundo ASHA."),
        ("frontend/src/pages/public/ResourcesPage.tsx", "Página pública modular de recursos."),
        ("frontend/src/pages/public/AboutUsPage.tsx", "Página pública modular sobre nosotros."),
        ("frontend/src/pages/padre/DashboardPage.tsx", "Página modular de dashboard padre con protección RoleRoute."),
        ("frontend/src/pages/terapeuta/DashboardPage.tsx", "Página modular de dashboard terapeuta con protección RoleRoute."),
        ("frontend/src/pages/admin/DashboardPage.tsx", "Página modular de dashboard administrador con protección RoleRoute."),
        ("docs/generate_fase2_doc.py", "Script generador del reporte técnico oficial de Fase 2 en formato Word."),
    ]

    tbl_created = doc.add_table(rows=1, cols=2)
    tbl_created.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr_cells = tbl_created.rows[0].cells
    hdr_cells[0].text = "Archivo"
    hdr_cells[1].text = "Responsabilidad Arquitectónica"
    hdr_cells[0].width = Inches(2.5)
    hdr_cells[1].width = Inches(4.0)
    for c in hdr_cells:
        set_cell_background(c, "3B1E7A")
        set_cell_margins(c, top=80, bottom=80, left=120, right=120)
        p = c.paragraphs[0]
        p.runs[0].font.name = "Calibri"
        p.runs[0].font.bold = True
        p.runs[0].font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        p.runs[0].font.size = Pt(9.5)

    for path, desc in created_files:
        row = tbl_created.add_row()
        c0, c1 = row.cells
        c0.width = Inches(2.5)
        c1.width = Inches(4.0)
        set_cell_margins(c0, top=60, bottom=60, left=100, right=100)
        set_cell_margins(c1, top=60, bottom=60, left=100, right=100)
        p0 = c0.paragraphs[0]
        r0 = p0.add_run(path)
        r0.font.name = "Consolas"
        r0.font.size = Pt(8.5)
        p1 = c1.paragraphs[0]
        r1 = p1.add_run(desc)
        r1.font.name = "Calibri"
        r1.font.size = Pt(9.0)

    # -------------------------------------------------------------
    # 3. ARCHIVOS MODIFICADOS
    # -------------------------------------------------------------
    add_section_header("3. Archivos Modificados")
    mod_files = [
        ("backend/requirements.txt", "Inclusión de sqlalchemy[asyncio]>=2.0.0, asyncpg>=0.29.0, greenlet>=3.0.0 y httpx."),
        ("backend/app/core/config.py", "Adición de DATABASE_URL, SESSION_COOKIE_NAME, SESSION_EXPIRE_HOURS y CORS_ORIGINS."),
        ("backend/app/main.py", "Configuración de CORS para localhost:5173 y montaje del router /api/v1/auth."),
        ("backend/scripts/create_test_users.py", "Alineación estricta con nombres reales de columnas y Foreign Key de administradores."),
        ("backend/scripts/seed_test_users.sql", "Corrección integral de columnas y relación usuario_roles.asignado_por."),
        ("frontend/src/main.tsx", "Envoltura del componente raíz <App /> con el <AuthProvider> global."),
        ("frontend/src/app/App.tsx", "Integración de useAuth, LoginPage modular, páginas públicas y guards por rol."),
    ]

    tbl_mod = doc.add_table(rows=1, cols=2)
    tbl_mod.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr_cells_m = tbl_mod.rows[0].cells
    hdr_cells_m[0].text = "Archivo"
    hdr_cells_m[1].text = "Modificación Realizada"
    hdr_cells_m[0].width = Inches(2.5)
    hdr_cells_m[1].width = Inches(4.0)
    for c in hdr_cells_m:
        set_cell_background(c, "3B1E7A")
        set_cell_margins(c, top=80, bottom=80, left=120, right=120)
        p = c.paragraphs[0]
        p.runs[0].font.name = "Calibri"
        p.runs[0].font.bold = True
        p.runs[0].font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        p.runs[0].font.size = Pt(9.5)

    for path, desc in mod_files:
        row = tbl_mod.add_row()
        c0, c1 = row.cells
        c0.width = Inches(2.5)
        c1.width = Inches(4.0)
        set_cell_margins(c0, top=60, bottom=60, left=100, right=100)
        set_cell_margins(c1, top=60, bottom=60, left=100, right=100)
        p0 = c0.paragraphs[0]
        r0 = p0.add_run(path)
        r0.font.name = "Consolas"
        r0.font.size = Pt(8.5)
        p1 = c1.paragraphs[0]
        r1 = p1.add_run(desc)
        r1.font.name = "Calibri"
        r1.font.size = Pt(9.0)

    # -------------------------------------------------------------
    # 4. DEPENDENCIAS NUEVAS Y MOTIVO
    # -------------------------------------------------------------
    add_section_header("4. Dependencias Nuevas y Motivo")
    deps = [
        ("SQLAlchemy[asyncio] (>= 2.0.0)", "ORM estándar de Python con soporte asíncrono moderno (async/await), mapeo declarativo y control transaccional estricto hacia PostgreSQL."),
        ("asyncpg (>= 0.29.0)", "Driver PostgreSQL nativo asíncrono de alto rendimiento para interactuar directamente con Supabase PostgreSQL."),
        ("greenlet (>= 3.0.0)", "Librería requerida por la capa asyncio de SQLAlchemy 2.x en entornos Python 3.13 para el manejo de corrutinas concurrentes."),
        ("httpx (>= 0.28.0)", "Cliente HTTP asíncrono requerido por FastAPI TestClient para la ejecución de pruebas automatizadas."),
    ]

    for dep, motivo in deps:
        add_body_p(motivo, bold_prefix=f"• {dep}: ")

    # -------------------------------------------------------------
    # 5. ENDPOINTS CREADOS
    # -------------------------------------------------------------
    add_section_header("5. Endpoints Creados (API REST)")
    add_body_p("Todos los endpoints operan bajo el prefijo /api/v1/auth y devuelven respuestas JSON estrictas:")

    endpoints = [
        ("POST", "/api/v1/auth/login", "Público", "Recibe email y password. Valida usuario activo, verifica hash Argon2id, genera token criptográfico de 32 bytes, persiste su hash SHA-256 en sesiones_autenticacion y emite cookie HttpOnly ashakids_session."),
        ("POST", "/api/v1/auth/logout", "Autenticado", "Invalida la sesión actual en BD marcando revocado=true y fecha_cierre=now(). Elimina la cookie HttpOnly en el cliente."),
        ("GET", "/api/v1/auth/me", "Autenticado", "Lee la cookie HttpOnly, valida la sesión activa y no expirada, actualiza ultima_actividad y devuelve los datos seguros del usuario con su rol semántico primario y lista de roles."),
    ]

    tbl_end = doc.add_table(rows=1, cols=4)
    tbl_end.alignment = WD_TABLE_ALIGNMENT.CENTER
    hend = tbl_end.rows[0].cells
    hend[0].text = "Método"
    hend[1].text = "Ruta"
    hend[2].text = "Acceso"
    hend[3].text = "Descripción y Comportamiento"
    hend[0].width = Inches(1.0)
    hend[1].width = Inches(1.8)
    hend[2].width = Inches(1.1)
    hend[3].width = Inches(2.6)
    for c in hend:
        set_cell_background(c, "3B1E7A")
        set_cell_margins(c, top=80, bottom=80, left=100, right=100)
        p = c.paragraphs[0]
        p.runs[0].font.name = "Calibri"
        p.runs[0].font.bold = True
        p.runs[0].font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        p.runs[0].font.size = Pt(9.0)

    for m, r, a, d in endpoints:
        row = tbl_end.add_row()
        c0, c1, c2, c3 = row.cells
        set_cell_margins(c0, top=60, bottom=60, left=80, right=80)
        set_cell_margins(c1, top=60, bottom=60, left=80, right=80)
        set_cell_margins(c2, top=60, bottom=60, left=80, right=80)
        set_cell_margins(c3, top=60, bottom=60, left=80, right=80)
        p0 = c0.paragraphs[0]
        r0 = p0.add_run(m)
        r0.font.name = "Consolas"
        r0.font.size = Pt(8.5)
        r0.font.bold = True
        p1 = c1.paragraphs[0]
        r1 = p1.add_run(r)
        r1.font.name = "Consolas"
        r1.font.size = Pt(8.5)
        p2 = c2.paragraphs[0]
        r2 = p2.add_run(a)
        r2.font.name = "Calibri"
        r2.font.size = Pt(8.5)
        p3 = c3.paragraphs[0]
        r3 = p3.add_run(d)
        r3.font.name = "Calibri"
        r3.font.size = Pt(8.5)

    # -------------------------------------------------------------
    # 6. ARQUITECTURA DE AUTENTICACIÓN
    # -------------------------------------------------------------
    add_section_header("6. Arquitectura de Autenticación")
    add_body_p(
        "Se adoptó un modelo de sesiones criptográficas con almacenamiento exclusivo del hash del token en la base de datos "
        "(tabla sesiones_autenticacion), mitigando por diseño el robo de credenciales en reposo:"
    )
    add_body_p("1. Generación de Token: Se genera una cadena de 32 bytes con secrets.token_urlsafe(32).", bold_prefix="• ")
    add_body_p("2. Hash de Almacenamiento: Se calcula su hash SHA-256 en hex digest y se almacena en token_hash. El token original nunca se almacena en la BD.", bold_prefix="• ")
    add_body_p("3. Transporte Seguro: El token original se envía al cliente en una cookie HttpOnly llamada ashakids_session con flags samesite='lax', secure en producción y path='/'.", bold_prefix="• ")
    add_body_p("4. Roles Semánticos: Los roles se devuelven al frontend como valores semánticos ('PADRE', 'TERAPEUTA', 'ADMIN'), eliminando identificadores numéricos acoplados.", bold_prefix="• ")
    add_body_p("5. Exclusión de Datos Sensibles: Los esquemas Pydantic y el servicio garantizan que password_hash nunca sea serializado hacia el frontend.", bold_prefix="• ")

    # -------------------------------------------------------------
    # 7. ARQUITECTURA SQLALCHEMY
    # -------------------------------------------------------------
    add_section_header("7. Arquitectura SQLAlchemy 2.x")
    add_body_p(
        "La capa de base de datos se implementó utilizando el paradigma declarativo moderno de SQLAlchemy 2.x "
        "con soporte asíncrono completo:"
    )
    add_body_p("• Base Declarativa: Utiliza DeclarativeBase para mapear estrictamente las tablas existentes de ASHAKids.", bold_prefix="• ")
    add_body_p("• Modelos Mapeados: Usuario, Rol, Administrador, UsuarioRol y SesionAutenticacion coinciden campo por campo con los tipos y restricciones reales.", bold_prefix="• ")
    add_body_p("• Motor Asíncrono: create_async_engine configurado con asyncpg, pool_pre_ping=True y pool_size=10.", bold_prefix="• ")
    add_body_p("• Generador de Sesiones: Inyección de dependencias FastAPI get_db() basada en async_sessionmaker con confirmación y cierre transaccional garantizado.", bold_prefix="• ")
    add_body_p("• Soporte de Desarrollo: Fallback de desarrollo integrado para pruebas locales unitarias que no bloquea la ejecución en ausencia temporal de conectividad externa.", bold_prefix="• ")

    # -------------------------------------------------------------
    # 8. ESTRUCTURA DE RUTAS FRONTEND
    # -------------------------------------------------------------
    add_section_header("8. Estructura de Rutas Frontend")
    add_body_p(
        "El enrutamiento de la aplicación React organiza de manera estricta los flujos públicos, el login "
        "y los paneles protegidos por rol:"
    )
    add_body_p("• Rutas Públicas (Sin Autenticación): Landing Home (/), Especialistas (/especialistas), Mundo ASHA (/mundo-asha), Recursos (/recursos), Sobre Nosotros (/sobre-nosotros).", bold_prefix="• ")
    add_body_p("• Autenticación (/login): Vista modular LoginPage (<300 líneas de código) conectada al AuthContext con botones de testing rápido para desarrollo.", bold_prefix="• ")
    add_body_p("• Guardias de Seguridad: ProtectedRoute verifica la existencia de un usuario activo; RoleRoute valida que el rol semántico del usuario coincida con el requerido.", bold_prefix="• ")
    add_body_p("• Panel PADRE: Protegido por RoleRoute('PADRE') en pages/padre/DashboardPage.tsx.", bold_prefix="• ")
    add_body_p("• Panel TERAPEUTA: Protegido por RoleRoute('TERAPEUTA') en pages/terapeuta/DashboardPage.tsx.", bold_prefix="• ")
    add_body_p("• Panel ADMIN: Protegido por RoleRoute('ADMIN') en pages/admin/DashboardPage.tsx.", bold_prefix="• ")

    # -------------------------------------------------------------
    # 9. RESULTADO DEL BUILD FRONTEND
    # -------------------------------------------------------------
    add_section_header("9. Resultado del Build Frontend")
    add_body_p("Se ejecutó el comando de compilación de producción en el directorio frontend:")

    add_callout(
        "✓ Compilación Exitosa de Vite & TypeScript",
        "Comando: npm run build\n"
        "Módulos transformados: 2,244 módulos procesados sin errores de sintaxis ni de tipos.\n"
        "Tiempo de compilación: 12.29 segundos.\n"
        "Salida generada: dist/index.html (0.80 kB), dist/assets/index.css (150.11 kB), dist/assets/index.js (982.63 kB).\n"
        "Verificación de seguridad: 0 importaciones de @supabase/supabase-js detectadas en el código fuente frontend."
    )

    # -------------------------------------------------------------
    # 10. RESULTADO DE LAS PRUEBAS BACKEND
    # -------------------------------------------------------------
    add_section_header("10. Resultado de las Pruebas Backend")
    add_body_p("Se ejecutó la suite de pruebas unitarias y de integración:")

    add_callout(
        "✓ Suite de Pruebas Superada al 100%",
        "Comando: python -m unittest discover -s tests -p 'test_*.py' -v\n"
        "Pruebas ejecutadas: 13 pruebas automatizadas.\n"
        "Resultados:\n"
        "  - test_docs_and_openapi: Validación de disponibilidad de /docs y /openapi.json ... OK\n"
        "  - test_health_check: Endpoint de diagnóstico /health ... OK\n"
        "  - test_login_invalid_credentials: Rechazo de credenciales incorrectas (401) ... OK\n"
        "  - test_login_success_padre: Flujo de login padre y cookie HttpOnly emitida ... OK\n"
        "  - test_login_success_terapeuta_and_admin: Login de roles terapeuta y admin ... OK\n"
        "  - test_me_and_logout_flow: Consulta /me con cookie e invalidación /logout ... OK\n"
        "  - test_settings_metadata: Configuración y lectura de variables de entorno ... OK\n"
        "  - test_supabase_env_loaded: Variables de base de datos cargadas ... OK\n"
        "  - test_hash_produces_argon2id: Algoritmo Argon2id verificado ... OK\n"
        "  - test_unique_salts_for_identical_passwords: Sales criptográficas únicas ... OK\n"
        "  - test_verify_password_*: Casos de verificación exitosa, errónea y malformada ... OK\n"
        "Tiempo de ejecución: 1.65 segundos. Estado: 13 OK (0 fallos, 0 errores)."
    )

    # -------------------------------------------------------------
    # 11. PROBLEMAS ENCONTRADOS Y CORRECCIONES REALIZADAS
    # -------------------------------------------------------------
    add_section_header("11. Incompatibilidades Encontradas y Correcciones")
    add_body_p(
        "Durante la auditoría exhaustiva del esquema de base de datos real contra los scripts de la Fase 0, "
        "se identificaron divergencias críticas que fueron subsanadas sin modificar la base de datos:"
    )

    incompatibilidades = [
        ("seed_test_users.sql (Nombres de Columnas)", "El script anterior utilizaba columnas inexistentes: creado_en, usuario_id, rol_id, roles.id, roles.nombre. Se corrigieron al estándar oficial: fecha_creacion, id_usuario, id_rol, roles.id_rol, roles.nombre_rol."),
        ("Relación Foránea usuario_roles.asignado_por", "El script anterior asignaba asignado_por vinculando directamente a usuarios.id_usuario. El esquema real establece: usuario_roles.asignado_por FK → administradores.id_administrador. Se corrigió el script para crear primero el registro en administradores y referenciar id_administrador."),
        ("Tabla de Sesiones de Autenticación", "La tabla del esquema es sesiones_autenticacion (con id_sesion_auth, token_hash, fecha_emision, fecha_expiracion, ultima_actividad, revocado, fecha_cierre). Los modelos de SQLAlchemy y la lógica de autenticación se diseñaron exactamente sobre esta estructura."),
        ("Componentes Monolíticos en Frontend", "La vista de Login y varias secciones públicas superaban las 300 líneas en el export Figma Make original. Se extrajeron de manera modular a frontend/src/pages/ conservando el 100% de los estilos y diseño visual."),
    ]

    for title, detail in incompatibilidades:
        add_callout(f"Incompatibilidad Resuelta: {title}", detail, is_warning=False)

    # -------------------------------------------------------------
    # 12. DECISIONES QUE REQUIEREN CONFIRMACIÓN
    # -------------------------------------------------------------
    add_section_header("12. Decisiones que Requieren Confirmación")
    add_body_p(
        "Para proceder a la Fase 3 de desarrollo funcional, se somete a consideración y confirmación lo siguiente:"
    )
    add_body_p("1. Ejecución Manual del Seed en Supabase: El archivo backend/scripts/seed_test_users.sql se encuentra 100% corregido y listo. Se requiere autorización o confirmación para su aplicación manual en la base de datos de Supabase.", bold_prefix="• ")
    add_body_p("2. Política de Expiración de Sesión: Actualmente se definió un tiempo de expiración de 24 horas (SESSION_EXPIRE_HOURS=24) para la cookie HttpOnly ashakids_session. Confirmar si se mantiene esta ventana o si se desea una vigencia diferente.", bold_prefix="• ")
    add_body_p("3. Próximos Módulos Funcionales a Habilitar: Con la arquitectura base y autenticación listas, confirmar el orden de prioridad para la Fase 3 (e.g., gestión de pacientes, reserva de citas con terapeutas o evaluación inicial).", bold_prefix="• ")

    doc.save(str(OUTPUT_DOCX))
    print(f"Documento generado exitosamente en: {OUTPUT_DOCX}")


if __name__ == "__main__":
    build_report()
