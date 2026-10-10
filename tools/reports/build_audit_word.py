"""Generate a new audit Word from a retained reference; preserve document chrome.

No database access. Body content comes from the versioned Markdown report.
"""
import argparse
from copy import deepcopy
from pathlib import Path
import re
from io import BytesIO
from zipfile import ZipFile, ZIP_DEFLATED

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt
from lxml import etree


def text_runs(paragraph, value, size=None):
    value = re.sub(r'\[([^\]]+)\]\(([^)]+)\)', r'\1 (\2)', value)
    for i, fragment in enumerate(re.split(r'\*\*(.*?)\*\*', value)):
        if not fragment:
            continue
        run = paragraph.add_run(fragment.replace('`', ''))
        run.bold = bool(i % 2)
        if size:
            run.font.size = Pt(size)


def build(source, template, output):
    document = Document(template)
    patterns = {name: deepcopy(next(p._p for p in document.paragraphs if p.style.name == name))
                for name in ('Title', 'Normal', 'Heading 1')}
    table_properties = deepcopy(document.tables[0]._tbl.tblPr)
    section = deepcopy(document.sections[0]._sectPr)
    for child in list(document._element.body):
        if child.tag != qn('w:sectPr'):
            document._element.body.remove(child)

    def paragraph(value, style='Normal'):
        p = document.add_paragraph(style=style)
        # Preserve paragraph-level layout from the actual reference pattern.
        source_pr = patterns[style].find(qn('w:pPr'))
        if source_pr is not None:
            if p._p.pPr is not None:
                p._p.remove(p._p.pPr)
            p._p.insert(0, deepcopy(source_pr))
        text_runs(p, value)
        if style != 'Normal':
            # Retained references use a bold navy title and black headings.
            first = patterns[style].find(qn('w:r'))
            original = first.find(qn('w:rPr')) if first is not None else None
            for run in p.runs:
                run.bold = True
                if original is not None:
                    if run._r.rPr is not None:
                        run._r.remove(run._r.rPr)
                    run._r.insert(0, deepcopy(original))
            p.paragraph_format.keep_with_next = True
        return p

    def table(lines):
        rows = [[s.strip() for s in line.strip('|').split('|')] for line in lines]
        rows = [r for r in rows if not all(re.fullmatch(r'[-: ]+', v) for v in r)]
        n = len(rows[0])
        ratios = [.28, .10, .10, .52] if n == 4 else ([.28, .42, .30] if n == 3 else [.32, .68])
        block = document.add_table(rows=len(rows), cols=n)
        block.autofit = False
        block._tbl.remove(block._tbl.tblPr)
        block._tbl.insert(0, deepcopy(table_properties))
        width = document.sections[0].page_width - document.sections[0].left_margin - document.sections[0].right_margin
        for col, ratio in zip(block.columns, ratios):
            col.width = int(width * ratio)
        for i, row in enumerate(block.rows):
            trpr = row._tr.get_or_add_trPr()
            if i == 0:
                repeat = OxmlElement('w:tblHeader'); trpr.append(repeat)
            no_split = OxmlElement('w:cantSplit'); trpr.append(no_split)
            for j, (cell, content) in enumerate(zip(row.cells, rows[i])):
                cell.width = int(width * ratios[j])
                cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                tcpr = cell._tc.get_or_add_tcPr()
                shade = OxmlElement('w:shd'); shade.set(qn('w:fill'), 'E7ECF2' if i == 0 else ('F5F8FB' if i % 2 == 0 else 'FFFFFF')); tcpr.append(shade)
                borders = OxmlElement('w:tcBorders')
                for side in ('top', 'left', 'bottom', 'right'):
                    border = OxmlElement(f'w:{side}'); border.set(qn('w:val'), 'single'); border.set(qn('w:sz'), '4'); border.set(qn('w:color'), 'D9D9D9'); borders.append(border)
                tcpr.append(borders)
                margins = OxmlElement('w:tcMar')
                for side in ('top', 'left', 'bottom', 'right'):
                    pad = OxmlElement(f'w:{side}'); pad.set(qn('w:w'), '70'); pad.set(qn('w:type'), 'dxa'); margins.append(pad)
                tcpr.append(margins)
                p = cell.paragraphs[0]; p.paragraph_format.space_after = Pt(4); p.paragraph_format.space_before = Pt(4)
                text_runs(p, content, 10)
                if i == 0:
                    for run in p.runs: run.bold = True
                if n == 4 and j in (1, 2): p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        document.add_paragraph().paragraph_format.space_after = Pt(2)

    lines = source.read_text(encoding='utf-8').splitlines()
    i = 0
    while i < len(lines):
        line = lines[i].strip(); i += 1
        if not line:
            continue
        if line == '<!-- pagebreak -->':
            document.add_page_break()
        elif line.startswith('# '):
            paragraph(line[2:], 'Title')
        elif line.startswith('## '):
            paragraph(line[3:], 'Heading 1')
        elif line.startswith('|'):
            group = [line]
            while i < len(lines) and lines[i].startswith('|'):
                group.append(lines[i]); i += 1
            table(group)
        elif line.startswith('!['):
            match = re.fullmatch(r'!\[([^\]]+)\]\(([^)]+)\)', line)
            if not match:
                raise ValueError('Invalid image')
            document.add_picture(str((source.parent / match.group(2)).resolve()), width=Inches(6.7))
            paragraph(match.group(1))
        else:
            paragraph(line)
    document.core_properties.title = lines[0].lstrip('# ')
    document.core_properties.subject = 'Auditoría técnica con evidencia y límites de alcance'
    document.core_properties.author = 'Equipo AshaKids | Evaluación técnica asistida'
    document.sections[0]._sectPr.getparent().replace(document.sections[0]._sectPr, section)
    update = OxmlElement('w:updateFields'); update.set(qn('w:val'), 'true'); document.settings.element.append(update)
    output.parent.mkdir(parents=True, exist_ok=True)
    draft = BytesIO()
    document.save(draft)
    with ZipFile(template) as original, ZipFile(draft) as generated:
        editable = {'word/document.xml', 'word/settings.xml', 'docProps/core.xml', 'word/_rels/document.xml.rels', '[Content_Types].xml'}
        names = set(original.namelist()) | set(generated.namelist())
        parts = {n: (generated.read(n) if n in editable or n.startswith('word/media/') else original.read(n))
                 for n in names if n in generated.namelist() or n not in editable}
        # Remove unused reference body images rather than carrying old evidence.
        relname = 'word/_rels/document.xml.rels'
        relationships = etree.fromstring(parts[relname]); body = etree.fromstring(parts['word/document.xml'])
        used = {value for el in body.iter() for key, value in el.attrib.items() if key.startswith('{http://schemas.openxmlformats.org/officeDocument/2006/relationships}')}
        for rel in list(relationships):
            if rel.get('Type', '').endswith('/image') and rel.get('Id') not in used:
                parts.pop('word/' + rel.get('Target'), None); relationships.remove(rel)
        parts[relname] = etree.tostring(relationships, xml_declaration=True, encoding='UTF-8', standalone=True)
        with ZipFile(output, 'w', ZIP_DEFLATED) as final:
            for name in sorted(parts): final.writestr(name, parts[name])
        for name in original.namelist():
            if name not in editable and not name.startswith('word/media/'):
                assert parts[name] == original.read(name), name
    print(output)


if __name__ == '__main__':
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--source', required=True, type=Path)
    parser.add_argument('--template', required=True, type=Path)
    parser.add_argument('--output', required=True, type=Path)
    args = parser.parse_args()
    build(args.source.resolve(), args.template.resolve(), args.output.resolve())
